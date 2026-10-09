"""Service xuất bảng lương Excel.
Cài đặt:  pip install fastapi uvicorn openpyxl supabase
Biến môi trường: SUPABASE_URL, SUPABASE_SERVICE_KEY
Chạy:     uvicorn main:app --reload
"""
import calendar, os
from collections import OrderedDict
from datetime import date, time
from decimal import Decimal
from io import BytesIO
from typing import Literal

from fastapi import FastAPI, Header, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THU = ["Hai", "Ba", "Tư", "Năm", "Sáu", "Bảy", "CN"]
COLS = ["Mã NV", "Tên nhân viên", "Chức vụ", "Ngày", "Thứ", "Vào 1", "Ra 1", "Vào 2", "Ra 2",
        "GIỜ CÔNG", "PHỤ CẤP CƠM", "LƯƠNG GIỜ", "LƯƠNG NGÀY", "LƯƠNG THÁNG",
        "TỔNG GIỜ CÔNG THÁNG", "NOTE", "TỔNG THU NHẬP", "SỐ TIỀN NOTE\n(dương = trừ, âm = cộng)"]
WIDTHS = [9, 22, 12, 12, 6, 8, 8, 8, 8, 10, 11, 11, 13, 14, 14, 38, 15, 16]
ORANGE = PatternFill("solid", fgColor="ED7D31")
LIGHT = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MERGE_COLS = (1, 2, 3, 14, 15, 17)  # Mã NV, Tên, Chức vụ, Lương tháng, Tổng giờ, Tổng thu nhập


def build_workbook(start: date, end: date, employees: list[dict]) -> Workbook:
    """employees: [{emp_code, full_name, position, hourly_rate,
                    days: [{date, in1, out1, in2, out2, note, adjust}]}]"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Bảng lương"
    font = lambda **k: Font(name="Arial", size=10, **k)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLS))
    t = ws.cell(1, 1, f"GIỜ CHẤM CÔNG (từ ngày {start:%d/%m/%Y} đến ngày {end:%d/%m/%Y})")
    t.font, t.fill = Font(name="Arial", size=14, bold=True, color="FFFFFF"), ORANGE
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    for i, h in enumerate(COLS, 1):
        c = ws.cell(3, i, h)
        c.font, c.fill, c.border = font(bold=True, color="FFFFFF"), ORANGE, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS[i - 1]
    ws.row_dimensions[3].height = 42
    ws.freeze_panes = "A4"

    r = 4
    for e in employees:
        days = e["days"]
        if not days:
            continue
        r0, r1 = r, r + len(days) - 1
        ws.cell(r0, 1, e["emp_code"]); ws.cell(r0, 2, e["full_name"]); ws.cell(r0, 3, e["position"])
        for d in days:
            ws.cell(r, 4, d["date"]).number_format = "DD/MM/YYYY"
            ws.cell(r, 5, THU[d["date"].weekday()])
            for col, key in zip((6, 7, 8, 9), ("in1", "out1", "in2", "out2")):
                ws.cell(r, col, d.get(key)).number_format = "HH:MM"
            # GIỜ CÔNG = (Ra1-Vào1)+(Ra2-Vào2); MOD(...,1) xử lý ca qua đêm
            ws.cell(r, 10, f'=IF(AND(F{r}<>"",G{r}<>""),MOD(G{r}-F{r},1)*24,0)'
                           f'+IF(AND(H{r}<>"",I{r}<>""),MOD(I{r}-H{r},1)*24,0)')
            ws.cell(r, 11, f'=IF(J{r}>=8,25000,0)')
            ws.cell(r, 12, e["hourly_rate"])
            ws.cell(r, 13, f"=J{r}*L{r}+K{r}")                      # LƯƠNG NGÀY
            ws.cell(r, 16, d.get("note"))
            ws.cell(r, 18, d.get("adjust") or 0)
            r += 1
        ws.cell(r0, 14, f"=SUM(M{r0}:M{r1})")                      # LƯƠNG THÁNG
        ws.cell(r0, 15, f"=SUM(J{r0}:J{r1})")                      # TỔNG GIỜ CÔNG THÁNG
        ws.cell(r0, 17, f"=N{r0}-SUM(R{r0}:R{r1})")               # TỔNG THU NHẬP

        for rr in range(r0, r1 + 1):
            for cc in range(1, len(COLS) + 1):
                c = ws.cell(rr, cc)
                c.border = BORDER
                c.font = font(bold=cc in (14, 17))
                c.alignment = Alignment(horizontal="left" if cc in (2, 16) else "center",
                                        vertical="center", wrap_text=cc == 16)
                if cc in (11, 12, 13, 14, 17, 18):
                    c.number_format = "#,##0"
                elif cc in (10, 15):
                    c.number_format = "0.00"
            ws.cell(rr, 17).fill = LIGHT
        if r1 > r0:
            for cc in MERGE_COLS:
                ws.merge_cells(start_row=r0, start_column=cc, end_row=r1, end_column=cc)
    return wb


# ---------- Pydantic: khớp từng cột của view v_payroll_days (01_schema.sql) ----------
class PayrollDayRow(BaseModel):
    """Kiểu dữ liệu tương ứng PostgreSQL -> JSON (PostgREST) -> Python:
    uuid -> str | numeric(12,0) -> Decimal | date -> date ('YYYY-MM-DD') | time -> time ('HH:MM:SS')"""
    user_id: str
    emp_code: str
    full_name: str
    position: str | None = None
    hourly_rate: Decimal
    allowance_per_shift: Decimal
    work_date: date
    check_in_1: time | None = None
    check_out_1: time | None = None
    check_in_2: time | None = None
    check_out_2: time | None = None
    note_text: str | None = None
    adjustment_amount: Decimal = Decimal(0)


# ---------- FastAPI ----------
app = FastAPI(title="Payroll export")
# Production: đặt CORS_ORIGINS="https://app-cua-ban.vercel.app" (nhiều domain ngăn cách bằng dấu phẩy)
app.add_middleware(CORSMiddleware, allow_origins=os.getenv("CORS_ORIGINS", "*").split(","),
                   allow_methods=["GET", "POST"], allow_headers=["*"])


def _sb():
    """Client Supabase dùng service_role (bỏ qua RLS) - chỉ chạy trên server."""
    from supabase import create_client
    return create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_SERVICE_KEY"])


def _require_manager(sb, authorization: str) -> str:
    """Xác thực JWT của người gọi và bắt buộc role 'manager'. Trả về uid."""
    try:
        uid = sb.auth.get_user(authorization.removeprefix("Bearer ").strip()).user.id
    except Exception:
        raise HTTPException(401, "Token không hợp lệ hoặc đã hết hạn")
    row = sb.table("users").select("role").eq("id", uid).maybe_single().execute()
    if not row or not row.data or row.data["role"] != "manager":
        raise HTTPException(403, "Chỉ quản lý được thực hiện thao tác này")
    return uid


def login_email(emp_code: str) -> str:
    """Quy ước email đăng nhập - phải trùng với hàm SQL resolve_login()."""
    return f"{emp_code.strip().lower()}@company.local"


@app.get("/health")
def health():
    return {"ok": True}


# /export-payroll là đường dẫn chính (theo yêu cầu); /payroll/export giữ lại để tương thích bản cũ
@app.get("/export-payroll")
@app.get("/payroll/export")
def export(month: int, year: int, authorization: str = Header(...)):
    if not 1 <= month <= 12:
        raise HTTPException(400, "Tháng phải từ 1 đến 12")
    sb = _sb()
    _require_manager(sb, authorization)

    start, end = date(year, month, 1), date(year, month, calendar.monthrange(year, month)[1])
    raw = (sb.table("v_payroll_days").select("*")
             .gte("work_date", start.isoformat()).lte("work_date", end.isoformat())
             .order("emp_code").order("work_date").execute().data)
    rows = [PayrollDayRow(**x) for x in raw]  # lỗi kiểu dữ liệu sẽ báo ngay tại đây

    emps: "OrderedDict[str, dict]" = OrderedDict()
    for x in rows:
        e = emps.setdefault(x.emp_code, dict(
            emp_code=x.emp_code, full_name=x.full_name, position=x.position,
            hourly_rate=int(x.hourly_rate), days=[]))
        e["days"].append(dict(date=x.work_date, in1=x.check_in_1, out1=x.check_out_1,
                              in2=x.check_in_2, out2=x.check_out_2,
                              note=x.note_text, adjust=int(x.adjustment_amount)))
    buf = BytesIO()
    build_workbook(start, end, list(emps.values())).save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="bang-luong-{year}-{month:02d}.xlsx"',
                 "Access-Control-Expose-Headers": "Content-Disposition"})


# ---------- Quản lý nhân viên (cần service_role nên đi qua backend) ----------
class NewUser(BaseModel):
    emp_code: str = Field(min_length=1, max_length=20, pattern=r"^[A-Za-z0-9_.-]+$")
    full_name: str = Field(min_length=1)
    password: str = Field(min_length=6)
    phone: str | None = None
    position: str | None = None
    role: Literal["employee", "manager"] = "employee"
    device_pin: str | None = None
    bank_name: str | None = None
    bank_account_no: str | None = None
    bank_account_name: str | None = None
    hourly_rate: int = Field(0, ge=0)
    allowance_per_shift: int = Field(0, ge=0)


class NewPassword(BaseModel):
    password: str = Field(min_length=6)


@app.post("/admin/users", status_code=201)
def create_user(body: NewUser, authorization: str = Header(...)):
    sb = _sb()
    _require_manager(sb, authorization)
    try:  # 1) tạo tài khoản đăng nhập
        uid = sb.auth.admin.create_user({"email": login_email(body.emp_code), "password": body.password,
                                          "email_confirm": True}).user.id
    except Exception as e:
        raise HTTPException(400, f"Không tạo được tài khoản (mã NV đã tồn tại?): {e}")
    try:  # 2) hồ sơ + lương; lỗi thì xoá tài khoản vừa tạo để không bị mồ côi
        data = body.model_dump(exclude={"password", "hourly_rate", "allowance_per_shift"})
        sb.table("users").insert({"id": uid, **data}).execute()
        sb.table("roles_rates").insert({"user_id": uid, "hourly_rate": body.hourly_rate,
                                        "allowance_per_shift": body.allowance_per_shift}).execute()
    except Exception as e:
        sb.auth.admin.delete_user(uid)
        raise HTTPException(400, f"Không lưu được hồ sơ (trùng SĐT/mã chấm công?): {e}")
    return {"id": uid}


@app.post("/admin/users/{uid}/password")
def reset_password(uid: str, body: NewPassword, authorization: str = Header(...)):
    sb = _sb()
    _require_manager(sb, authorization)
    try:
        sb.auth.admin.update_user_by_id(uid, {"password": body.password})
    except Exception as e:
        raise HTTPException(400, f"Không đổi được mật khẩu: {e}")
    return {"ok": True}
